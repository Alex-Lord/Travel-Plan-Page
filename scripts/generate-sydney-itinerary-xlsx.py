#!/usr/bin/env python3
"""Generate a concise bilingual, A4-print-ready Sydney itinerary workbook."""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "private" / "trip-data.json"
DESTINATION = ROOT / "Sydney-行程-A4打印版.xlsx"

NAVY = "17365D"
BLUE = "2C7DA0"
PALE_BLUE = "DCECF5"
PALE_SAND = "F8F1E7"
MUTED = "5E6B76"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="B7C7D2")

DAY_ENGLISH_TITLES = {
    1: "Arrival & Harbour Walk",
    2: "Taronga Zoo, Botanic Garden & The Rocks",
    3: "Bondi to Coogee Coastal Walk",
    4: "Blue Mountains Day Trip",
    5: "Manly, Shelly Beach & Darling Harbour",
    6: "Departure · Sydney to Beijing",
}

DAY_PLANS = {
    1: (
        "06:25 抵达；入境、寄存行李；海港步行",
        "Arrive 06:25; immigration, luggage drop; harbour walk",
        ["sydney-airport", "circular-quay", "opera-house", "botanic-garden", "the-rocks"],
        "若入境延误或疲劳，缩短 The Rocks / Observatory Hill。\nIf delayed or tired, skip The Rocks / Observatory Hill.",
    ),
    2: (
        "09:00 从 Song Hotel 出发；10:12 F2；Taronga Zoo；植物园与 The Rocks",
        "Leave Song Hotel 09:00; F2 at 10:12; Taronga Zoo; Botanic Garden & The Rocks",
        ["circular-quay", "taronga-zoo", "opera-house", "botanic-garden", "mrs-macquarie", "the-rocks"],
        "09:12 F2 不可行；09:32 只适合立即打车，乘火车/步行建议 10:12，并缩短下午步行。\nThe 09:12 F2 is not feasible; 09:32 requires an immediate taxi, while train/walk should target 10:12 and shorten the afternoon walk.",
    ),
    3: (
        "09:00 出发；大头贴；Bondi → Coogee 海岸徒步",
        "Leave 09:00; photobooth; Bondi → Coogee coastal walk",
        ["bondi-beach", "love-letters-photobooth", "tamarama", "bronte", "clovelly", "coogee"],
        "带水、防晒与防风层；体力不足可在 Bronte 结束。\nBring water, sun protection and a wind layer; finish at Bronte if needed.",
    ),
    4: (
        "07:30 BMT 火车；Echo Point；11:00 Scenic World；傍晚返悉尼",
        "07:30 BMT train; Echo Point; Scenic World at 11:00; return by evening",
        ["central-station", "katoomba", "echo-point", "scenic-world"],
        "只安排 Scenic World；雨雾或湿滑时及时缩短。\nScenic World only; shorten plans in fog, rain or wet conditions.",
    ),
    5: (
        "09:30 F1 轮渡；Manly 与 Shelly Beach；Darling Harbour 收尾",
        "09:30 F1 ferry; Manly & Shelly Beach; finish at Darling Harbour",
        ["circular-quay", "manly", "shelly-beach", "darling-harbour"],
        "前夜查 F1；最晚约 19:00 回酒店整理行李。\nCheck F1 the night before; return by about 19:00 to pack.",
    ),
    6: (
        "06:30 出发去机场；10:05 CZ326；广州转机返北京",
        "Leave for airport 06:30; CZ326 at 10:05; connect in Guangzhou to Beijing",
        ["central-station", "sydney-airport"],
        "优先保障国际航班；柜台确认两段登机牌与行李直挂。\nPrioritise the international flight; confirm both boarding passes and through-checked bags.",
    ),
}


def bilingual_place_name(place_id: str, places: dict[str, dict]) -> str:
    place = places.get(place_id, {})
    chinese = place.get("nameZh") or place.get("name") or place_id
    english = place.get("name") or chinese
    return chinese if chinese == english else f"{chinese} / {english}"


def configure_page(ws) -> None:
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.2, right=0.2, top=0.25, bottom=0.25, header=0.1, footer=0.1)
    ws.oddFooter.center.text = "Sydney Itinerary · 第 &P / &N 页"
    ws.oddFooter.center.size = 8
    ws.oddFooter.center.font = "Microsoft YaHei"
    ws.sheet_view.showGridLines = False


def style_cell(cell, *, fill: str | None = None, bold: bool = False, color: str = "17212B", size: float = 8.5, horizontal: str | None = None) -> None:
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(name="Microsoft YaHei", size=size, bold=bold, color=color)
    cell.alignment = Alignment(horizontal=horizontal, vertical="top", wrap_text=True)
    cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def main() -> None:
    data = json.loads(SOURCE.read_text(encoding="utf-8"))
    places = {place["id"]: place for place in data.get("places", [])}

    workbook = Workbook()
    workbook.properties.creator = "Travel Plan Page"
    workbook.properties.title = "Sydney Bilingual A4 Itinerary"
    ws = workbook.active
    ws.title = "A4 双语行程"

    ws.merge_cells("A1:E1")
    title = ws["A1"]
    title.value = "悉尼行程 · Sydney Itinerary"
    title.fill = PatternFill("solid", fgColor=NAVY)
    title.font = Font(name="Microsoft YaHei", size=18, bold=True, color=WHITE)
    title.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 31

    ws.merge_cells("A2:E2")
    subtitle = ws["A2"]
    subtitle.value = "2026.09.26 — 2026.10.01  |  2 人 / 2 travellers  |  A4 landscape print edition"
    subtitle.fill = PatternFill("solid", fgColor=PALE_BLUE)
    subtitle.font = Font(name="Microsoft YaHei", size=9, color=MUTED)
    subtitle.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 22

    headers = ["日期\nDate", "主题\nTheme", "时间与安排\nTime & plan", "路线\nRoute", "重点提醒\nKey note"]
    for column, value in enumerate(headers, 1):
        cell = ws.cell(4, column, value)
        style_cell(cell, fill=BLUE, bold=True, color=WHITE, size=9, horizontal="center")
    ws.row_dimensions[4].height = 28

    widths = [13, 28, 43, 53, 44]
    for column, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(column)].width = width

    for row, day in enumerate(data["days"], 5):
        number = day["day"]
        plan_cn, plan_en, route_ids, note = DAY_PLANS[number]
        date = datetime.strptime(day["date"], "%Y-%m-%d")
        weekdays_cn = "一二三四五六日"
        date_value = f"D{number} · {date:%m/%d} 周{weekdays_cn[date.weekday()]}\nDay {number} · {date:%a, %d %b}"
        theme = f"{day['title']}\n{DAY_ENGLISH_TITLES[number]}"
        plan = f"{plan_cn}\n{plan_en}"
        route = " → \n".join(bilingual_place_name(place_id, places) for place_id in route_ids)
        values = [date_value, theme, plan, route, note]
        for column, value in enumerate(values, 1):
            cell = ws.cell(row, column, value)
            style_cell(cell, fill=PALE_SAND if number % 2 else None, bold=column in (1, 2), color=NAVY if column in (1, 2) else "17212B", size=8.5)
        ws.row_dimensions[row].height = 78

    ws.merge_cells("A12:E12")
    footer = ws["A12"]
    footer.value = (
        "通用提醒 / General: 两人各自固定使用交通支付介质；蓝山、海岸线与 Manly 出发前复核天气、开放与班次。\n"
        "Use a separate, consistent payment method per traveller; recheck weather, opening hours and services before Blue Mountains, coastal and Manly days."
    )
    style_cell(footer, fill=PALE_BLUE, color=MUTED, size=8)
    ws.row_dimensions[12].height = 30
    ws.print_title_rows = "$1:$4"
    configure_page(ws)

    workbook.save(DESTINATION)
    print(DESTINATION)


if __name__ == "__main__":
    main()
